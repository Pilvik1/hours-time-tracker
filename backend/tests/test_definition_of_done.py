from datetime import UTC, datetime, timedelta
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook


class Clock:
    def __init__(self, value: datetime):
        self.value = value

    def __call__(self) -> datetime:
        return self.value


def test_required_end_to_end_flow(client: TestClient, monkeypatch):
    project = client.post("/api/projects", json={"name": "Northstar"}).json()
    clock = Clock(datetime(2026, 9, 29, 7, 0, tzinfo=UTC))
    monkeypatch.setattr("time_tracker.services.timer.utc_now", clock)

    started = client.post(
        "/api/timer/start",
        json={"description": "Research synthesis", "project_id": project["id"]},
    )
    assert started.status_code == 201
    client.patch("/api/timer", json={"comment": "Updated assumptions and sensitivity analysis"})
    clock.value += timedelta(hours=2, minutes=30)
    stopped = client.post("/api/timer/stop", json={})
    assert stopped.status_code == 200

    manual = client.post(
        "/api/time-entries",
        json={
            "description": "Initial title",
            "comment": "Initial comment",
            "project_id": project["id"],
            "date": "2026-09-30",
            "start_time": "09:00:00",
            "end_time": "10:00:00",
        },
    ).json()
    edited = client.patch(
        f"/api/time-entries/{manual['id']}",
        json={"description": "Updated title", "comment": "Updated comment"},
    )
    assert edited.json()["description"] == "Updated title"

    week = client.get("/api/time-entries", params={"week": "2026-09-29"}).json()
    assert len(week["items"]) == 2
    assert week["aggregates"]["total_seconds"] == 12600
    assert [item["seconds"] for item in week["aggregates"]["daily_totals"]] == [9000, 3600]

    export = client.get(
        "/api/exports/xlsx",
        params={"start_date": "2026-09-28", "end_date": "2026-10-04"},
    )
    workbook = load_workbook(BytesIO(export.content))
    sheet = workbook["Timesheet"]
    first = [sheet.cell(7, column).value for column in range(1, 10)]
    assert first[0].date().isoformat() == "2026-09-29"
    assert first[1].strftime("%H:%M") == "09:00"
    assert first[2].strftime("%H:%M") == "11:30"
    assert first[3].total_seconds() == 9000
    assert first[5] == "Northstar"
    assert first[7] == "Research synthesis"
    assert first[8] == "Updated assumptions and sensitivity analysis"
