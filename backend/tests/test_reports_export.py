from io import BytesIO

from conftest import create_manual_entry
from fastapi.testclient import TestClient
from openpyxl import load_workbook


def test_report_groups_hours(client: TestClient, project: dict):
    create_manual_entry(client, project["id"])
    create_manual_entry(
        client,
        project["id"],
        date="2026-09-30",
        start_time="12:00:00",
        end_time="13:00:00",
    )
    response = client.get(
        "/api/reports/summary",
        params={"start_date": "2026-09-29", "end_date": "2026-09-30"},
    )
    assert response.status_code == 200
    body = response.json()
    assert body["total_seconds"] == 12600
    assert body["decimal_hours"] == 3.5
    assert body["by_project"][0]["name"] == "Northstar"
    assert body["by_project"][0]["seconds"] == 12600


def test_excel_export_contains_titles_comments_and_totals(client: TestClient, project: dict):
    client.patch("/api/settings", json={"name": "Ada Example", "employer_name": "Example Company"})
    create_manual_entry(
        client,
        project["id"],
        description="Research synthesis",
        comment="Updated assumptions and sensitivity analysis",
    )
    create_manual_entry(
        client,
        project["id"],
        description="Management meeting",
        comment="Discussed forecast",
        start_time="12:00:00",
        end_time="13:00:00",
    )
    response = client.get(
        "/api/exports/xlsx",
        params={"start_date": "2026-09-29", "end_date": "2026-09-29"},
    )
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    workbook = load_workbook(BytesIO(response.content), data_only=False)
    assert workbook.sheetnames == ["Timesheet", "Summary"]
    sheet = workbook["Timesheet"]
    assert sheet["B1"].value == "Ada Example"
    assert sheet["B2"].value == "Example Company"
    assert [sheet.cell(6, column).value for column in range(1, 12)] == [
        "Date",
        "Start",
        "End",
        "Duration",
        "Decimal Hours",
        "Project",
        "Client",
        "Title",
        "Comment",
        "Tags",
        "Billable",
    ]
    assert sheet["F7"].value == "Northstar"
    assert sheet["H7"].value == "Research synthesis"
    assert sheet["I7"].value == "Updated assumptions and sensitivity analysis"
    assert sheet["D7"].value.total_seconds() == 9000
    assert sheet["E7"].value == 2.5
    total_row = next(row for row in range(1, sheet.max_row + 1) if sheet.cell(row, 1).value == "TOTAL HOURS")
    assert sheet.cell(total_row, 4).value.total_seconds() == 12600
    assert workbook["Summary"]["B2"].value.total_seconds() == 12600


def test_empty_excel_export_is_valid(client: TestClient):
    response = client.get(
        "/api/exports/xlsx",
        params={"start_date": "2026-01-01", "end_date": "2026-01-31"},
    )
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content))
    assert workbook["Timesheet"]["B4"].value.total_seconds() == 0


def test_excel_rounding_keeps_exact_and_reported_duration(client: TestClient):
    client.patch("/api/settings", json={"rounding_increment": 15})
    client.post(
        "/api/time-entries",
        json={"description": "Short task", "date": "2026-09-29", "duration_seconds": 601},
    )
    response = client.get(
        "/api/exports/xlsx",
        params={"start_date": "2026-09-29", "end_date": "2026-09-29"},
    )
    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook["Timesheet"]
    assert sheet["D7"].value.total_seconds() == 601
    assert sheet["F7"].value.total_seconds() == 900
    assert sheet["B4"].value.total_seconds() == 601
    assert sheet["D4"].value.total_seconds() == 900
    assert workbook["Summary"]["B2"].value.total_seconds() == 601
    assert workbook["Summary"]["C2"].value.total_seconds() == 900
