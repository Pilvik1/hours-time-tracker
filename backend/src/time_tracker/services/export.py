from collections import defaultdict
from datetime import date
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from time_tracker import models
from time_tracker.time_utils import as_utc, decimal_hours, round_seconds

BASE_HEADERS = [
    "Date",
    "Start",
    "End",
    "Duration",
    "Decimal Hours",
    "Project",
    "Client",
    "Title",
    "Comment",
    "Tags",
    "Billable",
]


def _duration_value(seconds: int) -> float:
    return seconds / 86400


def create_xlsx(
    entries: list[models.TimeEntry],
    settings: models.UserSettings,
    start_date: date,
    end_date: date,
) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Timesheet"

    use_rounding = settings.rounding_increment > 0
    headers = BASE_HEADERS[:5]
    if use_rounding:
        headers += ["Reported Duration", "Reported Decimal Hours"]
    headers += BASE_HEADERS[5:]
    total_exact = sum(entry.duration_seconds for entry in entries)
    total_reported = sum(round_seconds(entry.duration_seconds, settings.rounding_increment) for entry in entries)

    sheet["A1"] = "Employee:"
    sheet["B1"] = settings.name
    sheet["A2"] = "Employer:"
    sheet["B2"] = settings.employer_name
    sheet["A3"] = "Period:"
    sheet["B3"] = f"{start_date.isoformat()} — {end_date.isoformat()}"
    sheet["A4"] = "Exact total hours:"
    sheet["B4"] = _duration_value(total_exact)
    sheet["B4"].number_format = "[h]:mm"
    if use_rounding:
        sheet["C4"] = "Reported total hours:"
        sheet["C4"].font = Font(bold=True)
        sheet["D4"] = _duration_value(total_reported)
        sheet["D4"].number_format = "[h]:mm"
    for row in range(1, 5):
        sheet[f"A{row}"].font = Font(bold=True)

    header_row = 6
    for column, title in enumerate(headers, 1):
        cell = sheet.cell(row=header_row, column=column, value=title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="334155")
        cell.alignment = Alignment(horizontal="center")

    zone = ZoneInfo(settings.timezone)
    daily: dict[date, list[int]] = defaultdict(lambda: [0, 0])
    project_totals: dict[str, list[int]] = defaultdict(lambda: [0, 0])
    client_totals: dict[str, list[int]] = defaultdict(lambda: [0, 0])
    for row_number, entry in enumerate(entries, header_row + 1):
        reported = round_seconds(entry.duration_seconds, settings.rounding_increment)
        daily[entry.date][0] += entry.duration_seconds
        daily[entry.date][1] += reported
        project_name = entry.project.name if entry.project else "No project"
        client = entry.project.client if entry.project else None
        client_name = client.name if client else "No client"
        project_totals[project_name][0] += entry.duration_seconds
        project_totals[project_name][1] += reported
        client_totals[client_name][0] += entry.duration_seconds
        client_totals[client_name][1] += reported

        start = as_utc(entry.start_at).astimezone(zone).time().replace(tzinfo=None) if entry.start_at else None
        end = as_utc(entry.end_at).astimezone(zone).time().replace(tzinfo=None) if entry.end_at else None
        values = [
            entry.date,
            start,
            end,
            _duration_value(entry.duration_seconds),
            decimal_hours(entry.duration_seconds),
        ]
        if use_rounding:
            values += [_duration_value(reported), decimal_hours(reported)]
        values += [
            entry.project.name if entry.project else "",
            client.name if client else "",
            entry.description,
            entry.comment or "",
            ", ".join(sorted(tag.name for tag in entry.tags)),
            "Yes" if entry.billable else "No",
        ]
        for column, value in enumerate(values, 1):
            sheet.cell(row=row_number, column=column, value=value)
        sheet.cell(row=row_number, column=1).number_format = "dd.mm.yyyy"
        sheet.cell(row=row_number, column=2).number_format = "hh:mm"
        sheet.cell(row=row_number, column=3).number_format = "hh:mm"
        sheet.cell(row=row_number, column=4).number_format = "[h]:mm"
        sheet.cell(row=row_number, column=5).number_format = "0.00"
        if use_rounding:
            sheet.cell(row=row_number, column=6).number_format = "[h]:mm"
            sheet.cell(row=row_number, column=7).number_format = "0.00"
        comment_column = headers.index("Comment") + 1
        sheet.cell(row=row_number, column=comment_column).alignment = Alignment(wrap_text=True, vertical="top")

    last_data_row = header_row + len(entries)
    last_column = get_column_letter(len(headers))
    if entries:
        table = Table(displayName="TimesheetTable", ref=f"A{header_row}:{last_column}{last_data_row}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True
        )
        sheet.add_table(table)
    else:
        sheet.auto_filter.ref = f"A{header_row}:{last_column}{header_row}"
    sheet.freeze_panes = f"A{header_row + 1}"

    totals_start = last_data_row + 3
    sheet.cell(row=totals_start, column=1, value="Daily totals").font = Font(bold=True)
    for offset, (day, totals) in enumerate(sorted(daily.items()), 1):
        row = totals_start + offset
        sheet.cell(row=row, column=1, value=day).number_format = "dd.mm.yyyy"
        sheet.cell(row=row, column=4, value=_duration_value(totals[0])).number_format = "[h]:mm"
        if use_rounding:
            sheet.cell(row=row, column=6, value=_duration_value(totals[1])).number_format = "[h]:mm"
    total_row = totals_start + len(daily) + 2
    sheet.cell(row=total_row, column=1, value="TOTAL HOURS").font = Font(bold=True)
    sheet.cell(row=total_row, column=4, value=_duration_value(total_exact)).number_format = "[h]:mm"
    sheet.cell(row=total_row, column=4).font = Font(bold=True)
    if use_rounding:
        sheet.cell(row=total_row, column=6, value=_duration_value(total_reported)).number_format = "[h]:mm"
        sheet.cell(row=total_row, column=6).font = Font(bold=True)

    widths = {
        "Date": 13,
        "Start": 10,
        "End": 10,
        "Duration": 12,
        "Decimal Hours": 15,
        "Reported Duration": 18,
        "Reported Decimal Hours": 23,
        "Project": 20,
        "Client": 20,
        "Title": 28,
        "Comment": 48,
        "Tags": 24,
        "Billable": 11,
    }
    for index, header in enumerate(headers, 1):
        sheet.column_dimensions[get_column_letter(index)].width = widths[header]

    summary = workbook.create_sheet("Summary")
    summary["A1"] = "Timesheet Summary"
    summary["A1"].font = Font(bold=True, size=14)
    summary["A2"] = "Overall total"
    summary["B1"] = "Exact Hours"
    summary["B1"].font = Font(bold=True)
    summary["B2"] = _duration_value(total_exact)
    summary["B2"].number_format = "[h]:mm"
    if use_rounding:
        summary["C1"] = "Reported Hours"
        summary["C1"].font = Font(bold=True)
        summary["C2"] = _duration_value(total_reported)
        summary["C2"].number_format = "[h]:mm"

    sections = [
        ("Date", [(day.isoformat(), totals) for day, totals in sorted(daily.items())]),
        ("Project", sorted(project_totals.items())),
        ("Client", sorted(client_totals.items())),
    ]
    row = 4
    for title, values in sections:
        summary.cell(row=row, column=1, value=title).font = Font(bold=True, color="FFFFFF")
        summary.cell(row=row, column=2, value="Exact Hours").font = Font(bold=True, color="FFFFFF")
        summary.cell(row=row, column=1).fill = PatternFill("solid", fgColor="334155")
        summary.cell(row=row, column=2).fill = PatternFill("solid", fgColor="334155")
        if use_rounding:
            summary.cell(row=row, column=3, value="Reported Hours").font = Font(bold=True, color="FFFFFF")
            summary.cell(row=row, column=3).fill = PatternFill("solid", fgColor="334155")
        for name, totals in values:
            row += 1
            summary.cell(row=row, column=1, value=name)
            summary.cell(row=row, column=2, value=_duration_value(totals[0])).number_format = "[h]:mm"
            if use_rounding:
                summary.cell(row=row, column=3, value=_duration_value(totals[1])).number_format = "[h]:mm"
        row += 2
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 14
    summary.column_dimensions["C"].width = 16

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
